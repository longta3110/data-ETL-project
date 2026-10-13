"""Reproducible UCI retail validation with optional transactional PostgreSQL loading."""
import argparse
from collections import Counter
import csv
from datetime import datetime
from decimal import Decimal, InvalidOperation
import hashlib
import json
import shutil
from pathlib import Path
import urllib.request
import zipfile

URL = 'https://archive.ics.uci.edu/static/public/352/online+retail.zip'
HEADERS = ['InvoiceNo', 'StockCode', 'Description', 'Quantity', 'InvoiceDate',
           'UnitPrice', 'CustomerID', 'Country']
FIELDS = ['source_row', 'invoice_no', 'stock_code', 'description', 'quantity',
          'invoice_date', 'unit_price', 'customer_id', 'country', 'line_type']


def text(value):
    return '' if value is None else str(value).strip()


def validate(values):
    """Return normalized values, possibly overlapping rejection reasons, and warnings."""
    invoice, stock, description, quantity, date, price, customer, country = values
    reasons, warnings = [], []
    invoice, stock, country = text(invoice), text(stock), text(country)
    if not invoice or not stock or not country:
        reasons.append('missing_required_identifier')
    try:
        q = Decimal(str(quantity))
        if not q.is_finite() or q != q.to_integral_value() or q == 0:
            raise ValueError
        quantity = int(q)
    except (InvalidOperation, ValueError, OverflowError):
        reasons.append('invalid_quantity')
        quantity = 0
    try:
        price = Decimal(str(price))
        if not price.is_finite() or price <= 0:
            raise ValueError
    except (InvalidOperation, ValueError):
        reasons.append('nonpositive_or_invalid_price')
        price = Decimal(0)
    if not isinstance(date, datetime):
        reasons.append('invalid_date')
    cancellation = invoice.upper().startswith('C')
    if quantity and ((cancellation and quantity > 0) or (not cancellation and quantity < 0)):
        reasons.append('invoice_quantity_sign_mismatch')
    if not text(customer):
        warnings.append('missing_customer_id')
    if not text(description):
        warnings.append('missing_description')
    row = [invoice, stock, text(description), quantity,
           date.isoformat() if isinstance(date, datetime) else text(date),
           str(price), text(customer), country, 'cancellation' if cancellation else 'sale']
    return row, reasons, warnings


def download(directory):
    directory.mkdir(parents=True, exist_ok=True)
    target = directory / 'Online Retail.xlsx'
    if target.exists():
        return target
    archive = directory / 'source.zip'
    with urllib.request.urlopen(URL, timeout=60) as response, archive.open('wb') as out:
        shutil.copyfileobj(response, out)
    with zipfile.ZipFile(archive) as z:
        # Extract only the expected member, never arbitrary archive paths.
        with z.open('Online Retail.xlsx') as source, target.with_suffix('.tmp').open('wb') as out:
            shutil.copyfileobj(source, out)
    target.with_suffix('.tmp').replace(target)
    return target


def run(source, output):
    from openpyxl import load_workbook
    with source.open('rb') as stream:
        digest = hashlib.file_digest(stream, 'sha256').hexdigest()
    output.mkdir(parents=True, exist_ok=True)
    workbook = load_workbook(source, read_only=True, data_only=True)
    rows = workbook.active.iter_rows(values_only=True)
    if list(next(rows)) != HEADERS:
        workbook.close()
        raise ValueError('Unexpected source schema; refusing to process')
    counts = Counter(source_rows=0, accepted_rows=0, quarantined_rows=0)
    rejection_reasons, warnings = Counter(), Counter()
    seen = set()
    sales, cancellations = Decimal(0), Decimal(0)
    with (output / 'accepted.csv').open('w', newline='', encoding='utf-8') as good, \
         (output / 'quarantine.csv').open('w', newline='', encoding='utf-8') as bad:
        accepted, rejected = csv.writer(good), csv.writer(bad)
        accepted.writerow(FIELDS)
        rejected.writerow(['source_row', 'reasons', 'original_values_json'])
        for number, values in enumerate(rows, start=2):
            counts['source_rows'] += 1
            row, reasons, advisory = validate(values)
            warnings.update(advisory)
            # Exact normalized duplicates are ambiguous without a source line ID.
            # Preserve both rows and flag them; do not silently delete legitimate sales.
            signature = tuple(row)
            if signature in seen:
                warnings['duplicate_candidate'] += 1
            seen.add(signature)
            if reasons:
                counts['quarantined_rows'] += 1
                rejection_reasons.update(reasons)
                rejected.writerow([number, '|'.join(reasons), json.dumps(values, default=str)])
                continue
            counts['accepted_rows'] += 1
            accepted.writerow([number, *row])
            amount = row[3] * Decimal(row[5])
            if row[-1] == 'sale':
                sales += amount
            else:
                cancellations -= amount
    workbook.close()
    if counts['source_rows'] != counts['accepted_rows'] + counts['quarantined_rows']:
        raise ValueError('Source row reconciliation failed')
    report = {
        'dataset_id': digest, 'source_url': URL,
        'counts': dict(counts), 'rejection_reasons': dict(rejection_reasons),
        'warnings': dict(warnings), 'gross_sales_gbp': str(sales),
        'cancellation_value_gbp': str(cancellations),
        'net_transaction_value_gbp': str(sales - cancellations),
        'duplicate_policy': 'retain_and_flag',
        'note': 'Accepted transaction value, not audited revenue. No inventory data. '
                'Source timestamps have no timezone. Reason counts can overlap.'
    }
    (output / 'quality.json').write_text(json.dumps(report, indent=2), encoding='utf-8')
    return report


def load(output, report):
    import psycopg
    # Connection reads standard PG* environment variables, never a committed password.
    with psycopg.connect('') as conn:
        with conn.cursor() as cur:
            cur.execute('SELECT pg_advisory_xact_lock(352)')
            cur.execute((Path(__file__).parent / 'sql/schema.sql').read_text(encoding='utf-8'))
            cur.execute('INSERT INTO retail.dataset(dataset_id, source_url, row_count) '
                        'VALUES (%s,%s,%s) ON CONFLICT (dataset_id) DO NOTHING',
                        (report['dataset_id'], URL, report['counts']['source_rows']))
            cur.execute('CREATE TEMP TABLE incoming (LIKE retail.staging_line INCLUDING DEFAULTS) ON COMMIT DROP')
            with cur.copy('COPY incoming FROM STDIN') as copy:
                with (output / 'accepted.csv').open(encoding='utf-8', newline='') as f:
                    for r in csv.DictReader(f):
                        copy.write_row([report['dataset_id'], int(r['source_row']), r['invoice_no'],
                                        r['stock_code'], r['description'] or None, int(r['quantity']),
                                        datetime.fromisoformat(r['invoice_date']), Decimal(r['unit_price']),
                                        r['customer_id'] or None, r['country'], r['line_type']])
            cur.execute('INSERT INTO retail.staging_line SELECT * FROM incoming '
                        'ON CONFLICT (dataset_id, source_row) DO NOTHING')
            cur.execute('SELECT count(*), COALESCE(sum(quantity*unit_price),0) '
                        'FROM retail.staging_line WHERE dataset_id=%s', (report['dataset_id'],))
            count, amount = cur.fetchone()
            if count != report['counts']['accepted_rows'] or amount != Decimal(report['net_transaction_value_gbp']):
                raise ValueError('Database reconciliation failed; transaction rolled back')
            cur.execute((Path(__file__).parent / 'sql/analysis.sql').read_text(encoding='utf-8'),
                        {'dataset_id': report['dataset_id']})
            columns = [c.name for c in cur.description]
            results = cur.fetchall()
    with (output / 'monthly_analysis.csv').open('w', newline='', encoding='utf-8') as f:
        writer = csv.writer(f)
        writer.writerow(columns)
        writer.writerows(results)


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--input', type=Path)
    parser.add_argument('--output', type=Path, default=Path('reports'))
    parser.add_argument('--load', action='store_true')
    args = parser.parse_args()
    report = run(args.input or download(Path('data')), args.output)
    if args.load:
        load(args.output, report)
    print(json.dumps(report, indent=2))
